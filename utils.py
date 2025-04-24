# utils.py
"""
Utility functions for the AI Research Buddy application.

This module centralizes reusable logic for:
- API interactions (Google Gemini, Serper).
- Parsing content from various file formats.
- Generating plots from data.
- Helper functions for core application logic (e.g., domain classification, prompt formatting).
"""

import streamlit as st
import requests # For making HTTP requests to external APIs
import json     # For handling JSON data (API requests/responses)
import fitz     # PyMuPDF for PDF parsing
import docx    # python-docx for DOCX parsing
import pandas as pd # For data manipulation, especially CSV handling
import io       # For handling in-memory byte streams (plots, file uploads)
import random   # For adding randomness (e.g., chatbot behavior, joke styles)
import re       # For parsing AI plot suggestions
import google.generativeai as genai # Google Gemini API library
import matplotlib.pyplot as plt     # For generating static plots

# --- Constants ---
# API Endpoints
SERPER_API_URL = "https://google.serper.dev/search"       # For web text search
SERPER_IMAGES_API_URL = "https://google.serper.dev/images" # For image search

# Application-specific constants
VALID_DOMAINS = ["Physics", "Chemistry", "Mathematics", "Biology", "Engineering", "Computer Science", "General"] # Research domains
GEMINI_MODEL_NAME = "gemini-1.5-flash-latest" # Specific Gemini model to use

# --- !!! Plot Type Constants - Ensure these are saved !!! ---
# Base plot types supported for actual generation and AI suggestions
BASE_PLOT_TYPES = ["Line Plot", "Bar Chart", "Scatter Plot"]
# Plot types including the AI choice option for the UI dropdown
UI_PLOT_TYPES = ["Default (AI Choice)"] + BASE_PLOT_TYPES


# --- API Call Functions ---
def get_api_keys():
    """
    Retrieves and validates API keys for Google Gemini and Serper from Streamlit secrets.
    Also configures the Google Gemini client upon successful key retrieval.

    Returns:
        tuple: (google_api_key, serper_api_key). Values can be None if keys are missing or invalid.
    """
    google_key = None; serper_key = None
    try:
        google_key = st.secrets.get("GOOGLE_API_KEY"); serper_key = st.secrets.get("SERPER_API_KEY")
        if not google_key: st.warning("GOOGLE_API_KEY not found in secrets.")
        if not serper_key: st.warning("SERPER_API_KEY not found in secrets. Web/image search disabled.")
        if google_key:
             try: genai.configure(api_key=google_key)
             except Exception as e: st.error(f"Failed to configure Google Gemini client: {e}"); google_key = None
        return google_key, serper_key
    except FileNotFoundError: st.error(".streamlit/secrets.toml not found."); return None, None
    except Exception as e: st.error(f"Error accessing secrets.toml: {e}"); return google_key, serper_key

def call_gemini_api(prompt, api_key, max_output_tokens=4096, temperature=0.7):
    """
    Interacts with the Google Gemini API to generate text content based on a prompt.

    Args:
        prompt (str): The complete prompt string for the model.
        api_key (str): The Google API key (must be valid).
        max_output_tokens (int): The maximum number of tokens for the response.
        temperature (float): The sampling temperature (controls creativity/randomness).

    Returns:
        str or None: The generated text content from the model, or None if an error occurs.
    """
    if not api_key: st.error("Google API Key is missing for the Gemini call."); return None
    try:
        genai.configure(api_key=api_key)
        generation_config = genai.types.GenerationConfig(max_output_tokens=max_output_tokens, temperature=temperature)
        safety_settings=[ {"category": c, "threshold": "BLOCK_MEDIUM_AND_ABOVE"} for c in ["HARM_CATEGORY_HARASSMENT", "HARM_CATEGORY_HATE_SPEECH", "HARM_CATEGORY_SEXUALLY_EXPLICIT", "HARM_CATEGORY_DANGEROUS_CONTENT"]]
        model = genai.GenerativeModel(GEMINI_MODEL_NAME, generation_config=generation_config, safety_settings=safety_settings)
        response = model.generate_content(prompt)
        if not response.candidates:
             block_reason = getattr(response.prompt_feedback, 'block_reason', "Unknown")
             st.error(f"Request blocked by Gemini safety settings: {block_reason}"); return None
        if hasattr(response, 'text') and response.text: return response.text.strip()
        else:
             try: return response.candidates[0].content.parts[0].text.strip()
             except (IndexError, AttributeError, ValueError, TypeError) as e: st.error(f"Could not extract text from Gemini response structure: {e}"); return None
    except Exception as e: st.error(f"An unexpected error occurred calling Google Gemini API: {e}"); return None

def run_serper_search(query, api_key, max_results=5):
    """
    Fetches organic web search results from the Serper API.

    Args:
        query (str): The search term(s).
        api_key (str): The Serper API key.
        max_results (int): Maximum number of results to retrieve.

    Returns:
        tuple: (search_context, sources) where:
               - search_context (str | None): Formatted string of results for LLM context.
               - sources (list[dict]): List of {'title', 'link'} for citation.
    """
    if not api_key: st.error("Serper API Key is missing for web search."); return None, []
    payload = json.dumps({"q": query, "num": max_results}); headers = {'X-API-KEY': api_key, 'Content-Type': 'application/json'}
    try:
        response = requests.post(SERPER_API_URL, headers=headers, data=payload, timeout=30); response.raise_for_status()
        results = response.json(); search_results_text_parts = []; sources_list = []
        if "organic" in results and results["organic"]:
            for item in results["organic"][:max_results]:
                title = item.get("title", "N/A"); link = item.get("link", "#"); snippet = item.get("snippet", "N/A")
                search_results_text_parts.append(f"Title: {title}\nSnippet: {snippet}\nURL: {link}\n---")
                sources_list.append({"title": title, "link": link})
        if not sources_list: st.warning(f"No relevant web search results found via Serper for '{query}'."); return None, []
        search_context_str = "\n".join(search_results_text_parts)
        return search_context_str, sources_list
    except requests.exceptions.RequestException as e: st.error(f"Error calling Serper API (Text Search): {e}"); return None, []
    except json.JSONDecodeError: st.error(f"Failed to decode Serper API response: {response.text}"); return None, []
    except Exception as e: st.error(f"An unexpected error occurred during Serper text search: {e}"); return None, []

def search_images_serper(query, api_key):
    """
    Fetches image search results from Serper API and returns the URL of the first image.

    Args:
        query (str): The image search term(s).
        api_key (str): The Serper API key.

    Returns:
        str or None: The image URL string, or None if no results or an error occurs.
    """
    if not api_key: return None
    payload = json.dumps({"q": query}); headers = {'X-API-KEY': api_key, 'Content-Type': 'application/json'}
    try:
        response = requests.post(SERPER_IMAGES_API_URL, headers=headers, data=payload, timeout=30); response.raise_for_status()
        results = response.json()
        if "images" in results and results["images"]:
            image_url = results["images"][0].get("imageUrl")
            if image_url: return image_url
            else: st.warning(f"First image result for '{query}' has no URL."); return None
        else: st.info(f"No relevant images found via Serper for '{query}'."); return None
    except requests.exceptions.RequestException as e: st.error(f"Error calling Serper API (Image Search): {e}"); return None
    except json.JSONDecodeError: st.error(f"Failed to decode Serper Image API response: {response.text}"); return None
    except Exception as e: st.error(f"An unexpected error during Serper image search: {e}"); return None

# --- Content Parsing Functions ---
def parse_text_file(uploaded_file):
    """Parses content from an uploaded .txt file."""
    try: uploaded_file.seek(0); return uploaded_file.read().decode("utf-8")
    except Exception as e: st.error(f"Error reading text file: {e}"); return None
def parse_pdf_file(uploaded_file):
    """Parses text content from an uploaded .pdf file using PyMuPDF."""
    try:
        uploaded_file.seek(0); text = ""
        with fitz.open(stream=uploaded_file.read(), filetype="pdf") as doc: text = "".join(page.get_text() for page in doc)
        return text
    except Exception as e: st.error(f"Error reading PDF file: {e}"); return None
def parse_docx_file(uploaded_file):
    """Parses text content from an uploaded .docx file."""
    try:
        uploaded_file.seek(0); doc = docx.Document(io.BytesIO(uploaded_file.read()))
        text = "\n".join([para.text for para in doc.paragraphs]); return text
    except Exception as e: st.error(f"Error reading DOCX file: {e}"); return None
def parse_markdown_file(uploaded_file):
    """Parses content from an uploaded .md file (treated as text)."""
    return parse_text_file(uploaded_file)
def parse_csv_file_to_df(uploaded_file):
    """Parses a .csv file into a pandas DataFrame with basic validation."""
    try:
        uploaded_file.seek(0); df = pd.read_csv(uploaded_file)
        if df.empty: st.error("CSV file is empty."); return None
        if len(df.columns) == 0: st.error("CSV file has no columns."); return None
        return df
    except (pd.errors.EmptyDataError, pd.errors.ParserError) as e: st.error(f"Error reading CSV: {e}. Check format."); return None
    except Exception as e: st.error(f"Error reading CSV file: {e}"); return None
def parse_uploaded_file(uploaded_file, return_df_for_csv=False):
    """
    Master parser for uploaded files. Determines file type and calls appropriate parser.

    Args:
        uploaded_file (UploadedFile): Streamlit file object.
        return_df_for_csv (bool): If True and file is CSV, return DataFrame, else return string.

    Returns:
        str | pd.DataFrame | None: Parsed content or None on error/unsupported type.
    """
    if uploaded_file is None: return None
    file_name = uploaded_file.name.lower()
    if file_name.endswith(".csv") and return_df_for_csv: return parse_csv_file_to_df(uploaded_file)
    elif file_name.endswith(".txt"): return parse_text_file(uploaded_file)
    elif file_name.endswith(".pdf"): return parse_pdf_file(uploaded_file)
    elif file_name.endswith(".docx"): return parse_docx_file(uploaded_file)
    elif file_name.endswith(".md"): return parse_markdown_file(uploaded_file)
    elif file_name.endswith(".csv"):
        df = parse_csv_file_to_df(uploaded_file)
        return df.to_string(index=False) if df is not None else None
    else: st.error(f"Unsupported file type: {uploaded_file.name}. Supported: .txt, .pdf, .docx, .md, .csv"); return None

# --- Graph Generation Function ---
def generate_plot(df, x_col, y_col, plot_type):
    """
    Generates a plot (Line, Bar, Scatter) from DataFrame columns using Matplotlib.

    Args:
        df (pd.DataFrame): Data source.
        x_col (str): Column name for X-axis.
        y_col (str): Column name for Y-axis.
        plot_type (str): Type of plot (must be in BASE_PLOT_TYPES).

    Returns:
        bytes or None: PNG image bytes of the plot, or None if plotting fails.
    """
    if df is None or x_col not in df.columns or y_col not in df.columns: st.error("Invalid data or columns for plotting."); return None
    if plot_type not in BASE_PLOT_TYPES: st.error(f"Invalid plot type '{plot_type}' provided for generation."); return None
    fig = None
    try:
        df_plot = df[[x_col, y_col]].copy()
        if plot_type != "Bar Chart": df_plot[x_col] = pd.to_numeric(df_plot[x_col], errors='coerce')
        df_plot[y_col] = pd.to_numeric(df_plot[y_col], errors='coerce')
        df_plot.dropna(inplace=True)
        if df_plot.empty: st.error(f"No valid numeric data for plot columns after cleaning."); return None
        fig, ax = plt.subplots()
        if plot_type == "Line Plot": ax.plot(df_plot[x_col], df_plot[y_col])
        elif plot_type == "Bar Chart": ax.bar(df_plot[x_col], df_plot[y_col])
        elif plot_type == "Scatter Plot": ax.scatter(df_plot[x_col], df_plot[y_col])
        ax.set_xlabel(x_col); ax.set_ylabel(y_col); ax.set_title(f"{plot_type} of {y_col} vs {x_col}")
        plt.xticks(rotation=45, ha='right'); plt.tight_layout()
        buf = io.BytesIO(); plt.savefig(buf, format='png'); buf.seek(0)
        return buf.getvalue()
    except Exception as e: st.error(f"Error during plot generation: {e}"); return None
    finally:
        if fig is not None and plt.fignum_exists(fig.number): plt.close(fig)

# --- Agent Logic Helpers ---
def classify_query_domain(query, api_key):
    """
    Uses Gemini to predict the primary scientific domain of a query.

    Args:
        query (str): The user's query.
        api_key (str): Google API key.

    Returns:
        str: The predicted domain from VALID_DOMAINS, or "General".
    """
    if not api_key: return "General"
    prompt = f"""Analyze query domain. Choose best fit from: {', '.join(VALID_DOMAINS)}. Respond ONLY with domain name. Query: "{query}" Domain:"""
    try:
        response = call_gemini_api(prompt, api_key, max_output_tokens=30, temperature=0.1)
        if response:
            cleaned = response.strip().replace("*", "").replace("`", "").capitalize()
            if cleaned in VALID_DOMAINS: return cleaned
        st.warning(f"Domain classification response '{response}' unclear, using 'General'.")
        return "General"
    except Exception as e: st.warning(f"Domain classification error: {e}"); return "General"

def format_final_prompt(
    user_prompt, specialist_area="General", search_context=None, task_type="query",
    word_limit=None, content_to_analyze=None, additional_instructions=None, data_context=None,
    column_info=None, graph_guidance=None):
    """
    Constructs the appropriate prompt for the Gemini API based on the specified task and inputs.
    Includes specific structures for different research tasks and plot-related tasks.
    """
    instruction = "You are an AI Research Assistant. Provide comprehensive, accurate, and professionally formatted responses using Markdown."
    if specialist_area != "General": instruction += f" Focus on aspects relevant to {specialist_area}."

    if task_type == "extract":
        prompt = f"{instruction}\n**Task:** Extract Information\n**Field:** {specialist_area}\n**Request:** {user_prompt}"
        if search_context: prompt += f"\n**Web Context:**\n---\n{search_context}\n---\nSynthesize from context and knowledge."
        else: prompt += "\nAnswer from internal knowledge."
        if word_limit: prompt += f"\n\nConstraint: Approx {word_limit} words."
        else: prompt += f"\n\nConstraint: Detailed."
    elif task_type == "analyze":
        base_instr = " Analyze the following content."
        if additional_instructions and additional_instructions.strip(): instruction += f"{base_instr} Focus on: {additional_instructions}"
        else: instruction += f"{base_instr} Provide concise summary (key points/findings)."
        prompt = f"{instruction}\n**Task:** Content Analysis/Summarization\n**Content:**\n---\n{content_to_analyze}\n---\n**Analysis/Summary:**"
        if word_limit: prompt += f"\n\nConstraint: Approx {word_limit} words."
    elif task_type == "query":
        prompt = f"{instruction}\n**Task:** Answer Query\n**Query:** {user_prompt}"
        if search_context: prompt += f"\n**Web Context:**\n---\n{search_context}\n---\nUse context and knowledge."
        else: prompt += "\nAnswer from internal knowledge."
    elif task_type == "summarize_graph":
        instruction = "You are an AI Data Analyst Assistant."
        prompt = f"{instruction}\n**Task:** Summarize Key Takeaways from Data Context\n**Data Context:** {data_context}\n\n**Request:** List the most important key takeaways/trends as concise bullet points (- or *). Focus on objective observations.\n**Key Takeaways:**"
    elif task_type == "suggest_plot": # Task to ask AI for plot suggestions
        instruction = "You are an AI Data Visualization Assistant."
        prompt = f"""{instruction}
        **Task:** Suggest the best plot type and columns for visualizing the provided data.
        **Available Columns & Data Info:**\n{column_info}
        **User Guidance (Optional):** {graph_guidance if graph_guidance else 'None provided.'}
        **Supported Plot Types:** {', '.join(BASE_PLOT_TYPES)}
        **Request:** Based on the data and guidance, suggest the *single best* plot type from the supported list, and the most appropriate column names for X and Y axes.
        **Output Format:** Respond *only* in the following format (replace placeholders):
        Plot Type: [Suggested Plot Type from Supported List]
        X-Axis: [Suggested X Column Name from Available Columns]
        Y-Axis: [Suggested Y Column Name from Available Columns]
        """
    else: # Fallback
        prompt = f"{instruction}\n**User Query:** {user_prompt}"
        if search_context: prompt += f"\n**Web Context:**\n---\n{search_context}\n---"
    return prompt

def parse_plot_suggestion(suggestion_text):
    """
    Parses the Gemini response for plot suggestions using regex.

    Args:
        suggestion_text (str): The text response from Gemini.

    Returns:
        tuple (str, str, str) or None: (plot_type, x_col, y_col) if successful, else None.
    """
    if not suggestion_text: return None
    # Regex to capture required fields, ignoring case for keys
    match = re.search(r"Plot Type:\s*(.*?)\s*\n\s*X-Axis:\s*(.*?)\s*\n\s*Y-Axis:\s*(.*)", suggestion_text, re.IGNORECASE | re.DOTALL)
    if match:
        plot_type, x_col, y_col = match.group(1).strip(), match.group(2).strip(), match.group(3).strip()
        # Validate if the suggested plot type is one we can actually generate
        if plot_type in BASE_PLOT_TYPES:
            return plot_type, x_col, y_col
        else:
            st.warning(f"AI suggested an unsupported plot type: '{plot_type}'. Please select manually.")
            return None # Indicate failure due to unsupported type
    else:
        st.error(f"Could not parse AI plot suggestion format from response: '{suggestion_text[:100]}...'")
        return None # Indicate failure due to format mismatch